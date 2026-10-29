"""Checks on data/clean/security_jobs_monthly.csv.

1. Ratio check: the openings-to-applicants ratio MHLW prints must equal
   active openings / active applicants, rounded to 2 decimals.
2. Totals check: MHLW's own yearly figures (the average of the 12 months)
   must equal the average of the monthly values we extracted.
3. No missing months.
4. Overlap: April 2022 - March 2023 is published under both classifications;
   show how far apart they are.

Writes data/clean/checks_report_jobs.md.
"""
import csv
import os
import re
from collections import defaultdict

import openpyxl

ROOT = os.path.join(os.path.dirname(__file__), "..")
CLEAN = os.path.join(ROOT, "data", "clean", "security_jobs_monthly.csv")
REPORT = os.path.join(ROOT, "data", "clean", "checks_report_jobs.md")
RAW = os.path.join(ROOT, "data", "raw", "mhlw")


def main():
    rows = list(csv.DictReader(open(CLEAN, encoding="utf-8")))
    d = defaultdict(dict)
    for r in rows:
        d[(r["classification"], r["occupation"], r["measure"])][r["month"]] = float(r["value"])
    lines = ["# Checks report: job-openings data", "", f"Rows checked: {len(rows)}", ""]
    problems = 0

    # 1. ratio
    lines += ["## 1. Ratio = active openings / active applicants", ""]
    bad = []
    n = 0
    for (cl, occ, meas), s in d.items():
        if meas != "active_openings_ratio":
            continue
        op, ap = d[(cl, occ, "active_openings")], d[(cl, occ, "active_applicants")]
        for m, v in s.items():
            n += 1
            calc = op[m] / ap[m]
            if abs(calc - v) > 0.0051:
                bad.append(f"- {cl} {occ} {m}: printed {v}, calculated {calc:.3f}")
    problems += len(bad)
    lines += bad or [f"All {n} monthly ratios match (within rounding)."]
    lines.append("")

    # 2. yearly averages printed by MHLW
    lines += ["## 2. MHLW's yearly figures vs the average of our monthly values", ""]
    from importlib import import_module
    ext = import_module("03_extract_mhlw_jobs")
    bad, n = [], 0
    for fname, cl in ext.FILES:
        wb = openpyxl.load_workbook(os.path.join(RAW, fname), read_only=True, data_only=True)
        for idx, (meas, _) in ext.SHEETS.items():
            if meas == "active_openings_ratio":
                continue
            rs = list(wb[wb.sheetnames[idx]].iter_rows(values_only=True))
            for r in rs:
                label = str(r[0]).strip() if r[0] else ""
                if label not in ext.OCCUPATIONS:
                    continue
                occ = ext.OCCUPATIONS[label][0]
                for c in range(1, len(r)):
                    if rs[3][c] or r[c] is None:
                        continue
                    head = str(rs[1][c])
                    y = int(re.match(r"(\d{4})", head).group(1))
                    if "年度" in head:  # fiscal year April-March
                        months = [f"{y}-{m:02d}" for m in range(4, 13)] + [f"{y+1}-{m:02d}" for m in range(1, 4)]
                    else:
                        months = [f"{y}-{m:02d}" for m in range(1, 13)]
                    s = d[(cl, occ, meas)]
                    if not all(m in s for m in months):
                        continue  # year not fully covered in this file
                    n += 1
                    avg = sum(s[m] for m in months) / 12
                    if abs(avg - r[c]) > 1:
                        bad.append(f"- {fname} {occ} {meas} {head}: printed {r[c]}, average of months {avg:.1f}")
    problems += len(bad)
    lines += bad or [f"All {n} yearly figures match the average of their 12 months."]
    lines.append("")

    # 3. gaps
    lines += ["## 3. Missing months", ""]
    gaps = []
    for (cl, occ, meas), s in sorted(d.items()):
        ms = sorted(s)
        y, m = map(int, ms[0].split("-"))
        expected = []
        while f"{y}-{m:02d}" <= ms[-1]:
            expected.append(f"{y}-{m:02d}")
            y, m = (y + 1, 1) if m == 12 else (y, m + 1)
        miss = sorted(set(expected) - set(ms))
        if miss:
            gaps.append(f"- {cl} {occ} {meas}: {miss}")
        if (occ, meas) == ("security", "active_openings_ratio"):
            lines.append(f"- {cl}: {ms[0]} to {ms[-1]}")
    problems += len(gaps)
    lines += gaps or ["", "No missing months in any series."]
    lines.append("")

    # 4. overlap
    lines += ["## 4. Overlap of the two classifications (security occupations, ratio)", "",
              "| Month | 4th-revision MHLW classification | Japan Standard Occupational Classification |", "|---|---|---|"]
    old = d[(ext.FILES[0][1], "security", "active_openings_ratio")]
    new = d[(ext.FILES[1][1], "security", "active_openings_ratio")]
    for m in sorted(set(old) & set(new)):
        lines.append(f"| {m} | {old[m]:.2f} | {new[m]:.2f} |")
    lines.append("")

    open(REPORT, "w", encoding="utf-8").write("\n".join(lines))
    print(f"problems: {problems}; report written to {os.path.relpath(REPORT, ROOT)}")


if __name__ == "__main__":
    import sys
    sys.path.insert(0, os.path.dirname(__file__))
    main()
